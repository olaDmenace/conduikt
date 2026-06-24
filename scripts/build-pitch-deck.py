"""
Convert PITCH-DECK.md to PITCH-DECK.pptx.

Conservative renderer: title at top + content body. Tables and code blocks
are rendered as monospace text so they survive PowerPoint editing.

Run from project root:
    python scripts/build-pitch-deck.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent

# Optional argv override: build-pitch-deck.py INPUT.md OUTPUT.pptx
if len(sys.argv) >= 3:
    SRC = Path(sys.argv[1]).resolve()
    OUT = Path(sys.argv[2]).resolve()
else:
    SRC = ROOT / "PITCH-DECK.md"
    OUT = ROOT / "PITCH-DECK.pptx"

TEAL = RGBColor(0x1F, 0x6B, 0x66)
COPPER = RGBColor(0xD9, 0x66, 0x3A)
OBSIDIAN = RGBColor(0x1A, 0x1A, 0x1A)
TEXT = RGBColor(0x2E, 0x2A, 0x24)
MUTED = RGBColor(0x6B, 0x66, 0x5E)
LIGHT_BG = RGBColor(0xF7, 0xF3, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def strip_frontmatter(text: str) -> str:
    lines = text.splitlines()
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                return "\n".join(lines[i + 1:])
    return text


def split_slides(text: str) -> list[str]:
    out, buf = [], []
    for line in text.splitlines():
        if line.strip() == "---":
            if buf:
                out.append("\n".join(buf).strip())
                buf = []
        else:
            buf.append(line)
    if buf:
        out.append("\n".join(buf).strip())
    return [s for s in out if s]


def strip_md(text: str) -> str:
    """Strip markdown emphasis markers from a line so it renders cleanly."""
    text = re.sub(r"\*\*\*(.+?)\*\*\*", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text


def is_table_row(line: str) -> bool:
    return line.count("|") >= 2


def is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line))


def render_table_as_text(lines: list[str]) -> str:
    """Convert markdown table lines into plain monospace text columns."""
    rows: list[list[str]] = []
    for ln in lines:
        if is_table_sep(ln):
            continue
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        cells = [strip_md(c) for c in cells]
        rows.append(cells)
    if not rows:
        return ""
    cols = max(len(r) for r in rows)
    widths = [0] * cols
    for r in rows:
        for i, c in enumerate(r):
            widths[i] = max(widths[i], len(c))
    out_lines = []
    for ri, r in enumerate(rows):
        padded = [r[i].ljust(widths[i]) if i < len(r) else " " * widths[i] for i in range(cols)]
        out_lines.append("  ".join(padded))
        if ri == 0:
            out_lines.append("  ".join("-" * widths[i] for i in range(cols)))
    return "\n".join(out_lines)


def parse_slide(raw: str) -> dict:
    """Return {title, subtitle, blocks: [{type, content}]} for one slide."""
    lines = raw.splitlines()
    title = None
    subtitle = None
    blocks: list[dict] = []

    i = 0
    # Title: first # / ## / ### at the top
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("## "):
            title = strip_md(s[3:])
            i += 1
            break
        if s.startswith("# "):
            title = strip_md(s[2:])
            i += 1
            break
        if s.startswith("### "):
            # Treat the very first ### as a subtitle only if no title yet
            if title is None:
                subtitle = strip_md(s[4:])
                i += 1
                continue
            break
        break

    # Optional subtitle right after title (### form)
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("### ") and subtitle is None:
            subtitle = strip_md(s[4:])
            i += 1
            continue
        break

    # Body
    buf_para: list[str] = []

    def flush_para():
        if buf_para:
            blocks.append({"type": "para", "content": " ".join(buf_para).strip()})
            buf_para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if not s:
            flush_para()
            i += 1
            continue

        if s.startswith("```"):
            flush_para()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # consume closing ```
            blocks.append({"type": "code", "content": "\n".join(code_lines)})
            continue

        if is_table_row(s) and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            flush_para()
            tbl_lines = []
            while i < len(lines) and is_table_row(lines[i].strip()):
                tbl_lines.append(lines[i])
                i += 1
            blocks.append({"type": "code", "content": render_table_as_text(tbl_lines)})
            continue

        if s.startswith(">"):
            flush_para()
            q_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q_lines.append(strip_md(lines[i].strip().lstrip(">").strip()))
                i += 1
            blocks.append({"type": "quote", "content": " ".join(q_lines)})
            continue

        if s.startswith("### "):
            flush_para()
            blocks.append({"type": "h3", "content": strip_md(s[4:])})
            i += 1
            continue
        if s.startswith("#### "):
            flush_para()
            blocks.append({"type": "h4", "content": strip_md(s[5:])})
            i += 1
            continue

        m = re.match(r"^\s*[-*]\s+(.+)$", line)
        if m:
            flush_para()
            items = []
            while i < len(lines):
                m2 = re.match(r"^\s*[-*]\s+(.+)$", lines[i])
                if not m2:
                    break
                items.append(strip_md(m2.group(1)))
                i += 1
            blocks.append({"type": "bullets", "content": items})
            continue

        m = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if m:
            flush_para()
            items = []
            while i < len(lines):
                m2 = re.match(r"^\s*\d+\.\s+(.+)$", lines[i])
                if not m2:
                    break
                items.append(strip_md(m2.group(1)))
                i += 1
            blocks.append({"type": "numbered", "content": items})
            continue

        buf_para.append(strip_md(s))
        i += 1

    flush_para()
    return {"title": title, "subtitle": subtitle, "blocks": blocks}


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.3), Inches(13.33), Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COPPER
    bar.line.fill.background()


def add_title_block(slide, title: str | None, subtitle: str | None) -> float:
    """Returns the y-position (in inches) where body content should start."""
    top = 0.4
    if not title and not subtitle:
        return 0.6
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    if title:
        p = tf.paragraphs[0]
        p.text = title
        r = p.runs[0]
        r.font.size = Pt(30)
        r.font.bold = True
        r.font.color.rgb = TEAL
    if subtitle:
        p = tf.add_paragraph() if title else tf.paragraphs[0]
        p.text = subtitle
        r = p.runs[0]
        r.font.size = Pt(16)
        r.font.italic = True
        r.font.color.rgb = MUTED
    return 1.6 if title else 0.6


def render_blocks(slide, blocks: list[dict], start_top: float):
    """Stack blocks into a single text frame for predictable layout."""
    body = slide.shapes.add_textbox(
        Inches(0.5), Inches(start_top), Inches(12.3), Inches(7.5 - start_top - 0.3)
    )
    tf = body.text_frame
    tf.word_wrap = True
    first = True

    def new_para():
        nonlocal first
        if first:
            first = False
            return tf.paragraphs[0]
        return tf.add_paragraph()

    for block in blocks:
        btype = block["type"]
        content = block["content"]

        if btype == "para":
            p = new_para()
            p.text = content
            r = p.runs[0]
            r.font.size = Pt(15)
            r.font.color.rgb = TEXT
            p.space_after = Pt(4)

        elif btype == "h3":
            p = new_para()
            p.text = content
            r = p.runs[0]
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = COPPER
            p.space_before = Pt(8)
            p.space_after = Pt(4)

        elif btype == "h4":
            p = new_para()
            p.text = content
            r = p.runs[0]
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = TEAL
            p.space_before = Pt(6)
            p.space_after = Pt(2)

        elif btype == "quote":
            p = new_para()
            p.text = "❝  " + content
            r = p.runs[0]
            r.font.size = Pt(16)
            r.font.italic = True
            r.font.color.rgb = TEAL
            p.space_before = Pt(6)
            p.space_after = Pt(6)

        elif btype == "code":
            for line in content.splitlines() or [""]:
                p = new_para()
                p.text = line if line else " "
                r = p.runs[0]
                r.font.name = "Consolas"
                r.font.size = Pt(11)
                r.font.color.rgb = TEXT

        elif btype == "bullets":
            for item in content:
                p = new_para()
                p.text = "▸  " + item
                r = p.runs[0]
                r.font.size = Pt(14)
                r.font.color.rgb = TEXT
                p.space_after = Pt(2)

        elif btype == "numbered":
            for idx, item in enumerate(content, 1):
                p = new_para()
                p.text = f"{idx}.  {item}"
                r = p.runs[0]
                r.font.size = Pt(14)
                r.font.color.rgb = TEXT
                p.space_after = Pt(2)


def main():
    print("Reading", SRC, flush=True)
    text = SRC.read_text(encoding="utf-8")
    text = strip_frontmatter(text)
    slides_raw = split_slides(text)
    print(f"Found {len(slides_raw)} slides", flush=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, raw in enumerate(slides_raw, 1):
        print(f"  rendering slide {idx}...", flush=True)
        slide = prs.slides.add_slide(blank)
        add_accent_bar(slide)
        parsed = parse_slide(raw)
        body_top = add_title_block(slide, parsed["title"], parsed["subtitle"])
        render_blocks(slide, parsed["blocks"], body_top)

    print("Saving", OUT, flush=True)
    prs.save(OUT)
    print(f"OK — wrote {OUT}", flush=True)


if __name__ == "__main__":
    sys.stdout.reconfigure(line_buffering=True)
    main()
